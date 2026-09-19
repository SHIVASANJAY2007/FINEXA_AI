"""
convert_reports.py

Converts Markdown reports into:
  1. Professional Microsoft Word (.docx) documents with styling, headings, and tables.
  2. High-quality PDF (.pdf) documents via Chrome/Edge headless rendering of styled HTML.
"""

import os
import re
import subprocess
import markdown
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
BROWSER_EXE = CHROME_PATH if os.path.exists(CHROME_PATH) else EDGE_PATH

WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS_DIR = os.path.join(WORKSPACE_DIR, "Reports")
FILES = [
    "Finexa_AI_Modules_Description_Report",
    "GWO_Technical_Report",
    "GWO_Evaluation_and_Accuracy_Report"
]

HTML_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800;900&family=Source+Serif+4:ital,opsz,wght@0,8..60,400..800;1,8..60,400..800&family=JetBrains+Mono:wght@400;600;700&display=swap');
  
  @page {{
    size: A4;
    margin: 18mm 16mm 18mm 16mm;
  }}

  body {{
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    color: #3A2E25;
    background: #FFFFFF;
    line-height: 1.6;
    font-size: 10.5pt;
    margin: 0;
    padding: 0;
  }}

  .header-banner {{
    border-bottom: 3px solid #0B4F4A;
    padding-bottom: 12px;
    margin-bottom: 20px;
    display: flex;
    justify-content: space-between;
    align-items: center;
  }}

  .header-logo {{
    font-size: 20pt;
    font-weight: 900;
    font-family: 'Source Serif 4', Georgia, serif;
    color: #3A2E25;
    letter-spacing: -0.5px;
  }}
  .header-logo span {{
    color: #C9A227;
  }}
  .header-tag {{
    font-size: 8.5pt;
    font-weight: 700;
    color: #0B4F4A;
    background: #E8F4F2;
    padding: 4px 10px;
    border-radius: 6px;
    text-transform: uppercase;
    letter-spacing: 1px;
    font-family: 'JetBrains Mono', monospace;
  }}

  h1 {{
    font-family: 'Source Serif 4', Georgia, serif;
    color: #3A2E25;
    font-size: 18pt;
    font-weight: 900;
    line-height: 1.25;
    margin-top: 20px;
    margin-bottom: 10px;
    padding-bottom: 6px;
    border-bottom: 1.5px solid #DCCFC0;
  }}

  h2 {{
    font-family: 'Source Serif 4', Georgia, serif;
    color: #0B4F4A;
    font-size: 13pt;
    font-weight: 800;
    margin-top: 18px;
    margin-bottom: 8px;
    padding-left: 8px;
    border-left: 4px solid #0B4F4A;
  }}

  h3 {{
    font-size: 11pt;
    font-weight: 750;
    color: #6B1E2B;
    margin-top: 15px;
    margin-bottom: 6px;
  }}

  h4 {{
    font-size: 10pt;
    font-weight: 700;
    color: #9A8678;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-top: 12px;
    margin-bottom: 4px;
  }}

  p {{
    margin: 6px 0;
    text-align: justify;
  }}

  blockquote {{
    background: #FDF6ED;
    border-left: 4px solid #C9A227;
    margin: 14px 0;
    padding: 10px 16px;
    border-radius: 0 8px 8px 0;
    color: #3A2E25;
    font-size: 10pt;
  }}

  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 14px 0;
    font-size: 9pt;
    page-break-inside: avoid;
  }}

  th, td {{
    border: 1px solid #DCCFC0;
    padding: 7px 10px;
    text-align: left;
    vertical-align: top;
  }}

  th {{
    background: #F6F3EB;
    color: #3A2E25;
    font-weight: 700;
    text-transform: uppercase;
    font-size: 8pt;
    letter-spacing: 0.5px;
    border-bottom: 2px solid #0B4F4A;
  }}

  tr:nth-child(even) {{
    background: #FDFBFA;
  }}

  code {{
    font-family: 'JetBrains Mono', Consolas, Monaco, monospace;
    font-size: 8.5pt;
    background: #F6F3EB;
    color: #0B4F4A;
    padding: 2px 5px;
    border-radius: 4px;
    border: 1px solid #E5DEC9;
  }}

  pre {{
    background: #F8F6F0;
    border: 1px solid #DCCFC0;
    border-radius: 8px;
    padding: 12px;
    overflow-x: auto;
    font-family: 'JetBrains Mono', Consolas, monospace;
    font-size: 8pt;
    line-height: 1.4;
    color: #3A2E25;
    margin: 12px 0;
  }}

  pre code {{
    background: transparent;
    border: none;
    padding: 0;
    color: inherit;
  }}

  ul, ol {{
    padding-left: 20px;
    margin: 6px 0;
  }}

  li {{
    margin-bottom: 3px;
  }}

  hr {{
    border: 0;
    border-top: 1px solid #DCCFC0;
    margin: 18px 0;
  }}

  .footer {{
    margin-top: 30px;
    padding-top: 12px;
    border-top: 1px solid #DCCFC0;
    font-size: 8pt;
    color: #9A8678;
    text-align: center;
  }}
</style>
</head>
<body>
  <div class="header-banner">
    <div class="header-logo">FINEXA <span>AI</span></div>
    <div class="header-tag">Official Algorithm & Architecture Report</div>
  </div>
  {content}
  <div class="footer">
    Generated by FINEXA AI Machine Learning Systems • Confidential Technical Report
  </div>
</body>
</html>
"""


def convert_to_html_and_pdf(md_path, html_path, pdf_path, title):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    body_html = markdown.markdown(md_text, extensions=['tables', 'fenced_code', 'toc', 'sane_lists'])
    full_html = HTML_TEMPLATE.format(title=title, content=body_html)

    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(full_html)

    print(f"[*] HTML generated: {html_path}")

    cmd = [
        BROWSER_EXE,
        "--headless=new",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        html_path
    ]
    subprocess.run(cmd, check=True)
    print(f"[+] PDF generated: {pdf_path}")


def convert_to_docx(md_path, docx_path, title):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(58, 46, 37)

    header_p = doc.add_paragraph()
    run_logo = header_p.add_run("FINEXA ")
    run_logo.bold = True
    run_logo.font.size = Pt(18)
    run_logo.font.color.rgb = RGBColor(58, 46, 37)

    run_ai = header_p.add_run("AI — ")
    run_ai.bold = True
    run_ai.font.size = Pt(18)
    run_ai.font.color.rgb = RGBColor(201, 162, 39)

    run_title = header_p.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.color.rgb = RGBColor(11, 79, 74)

    p_div = doc.add_paragraph()
    p_div_run = p_div.add_run("―" * 45)
    p_div_run.font.color.rgb = RGBColor(220, 207, 192)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for l in t_lines:
            cells = [c.strip() for c in l.strip().strip('|').split('|')]
            if all(set(c).issubset({'-', ':', ' '}) for c in cells):
                continue
            rows_data.append(cells)
        
        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, r_data in enumerate(rows_data):
            row = table.rows[row_idx]
            for col_idx in range(cols_count):
                cell = row.cells[col_idx]
                val = r_data[col_idx] if col_idx < len(r_data) else ""
                cell.text = val
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.space_after = Pt(2)
                cell_p.paragraph_format.space_before = Pt(2)
                
                if row_idx == 0:
                    shading_elm = parse_xml(r'<w:shd {} w:fill="0B4F4A"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shading_elm)
                    for run in cell_p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(253, 246, 237)
                elif row_idx % 2 == 1:
                    shading_elm = parse_xml(r'<w:shd {} w:fill="F6F3EB"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shading_elm)

        doc.add_paragraph()

    for line in lines:
        raw = line.rstrip('\r\n')
        stripped = raw.strip()

        if stripped.startswith('```'):
            if in_code_block:
                p_code = doc.add_paragraph()
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(4)
                run_code = p_code.add_run('\n'.join(code_lines))
                run_code.font.name = 'Consolas'
                run_code.font.size = Pt(9.5)
                run_code.font.color.rgb = RGBColor(11, 79, 74)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw)
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not stripped:
            continue

        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            r = h.add_run(stripped[2:])
            r.font.size = Pt(16)
            r.font.color.rgb = RGBColor(58, 46, 37)
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            r = h.add_run(stripped[3:])
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(11, 79, 74)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            r = h.add_run(stripped[4:])
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(107, 30, 43)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run(stripped[2:])
            r.italic = True
            r.font.color.rgb = RGBColor(90, 75, 60)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"[+] DOCX generated: {docx_path}")


if __name__ == '__main__':
    if not os.path.exists(REPORTS_DIR):
        os.makedirs(REPORTS_DIR, exist_ok=True)

    for name in FILES:
        md_file = os.path.join(REPORTS_DIR, f"{name}.md")
        html_file = os.path.join(REPORTS_DIR, f"{name}.html")
        pdf_file = os.path.join(REPORTS_DIR, f"{name}.pdf")
        docx_file = os.path.join(REPORTS_DIR, f"{name}.docx")
        
        title_str = name.replace('_', ' ')
        print(f"\nProcessing {name} in Reports/...")
        if os.path.exists(md_file):
            convert_to_html_and_pdf(md_file, html_file, pdf_file, title_str)
            convert_to_docx(md_file, docx_file, title_str)
        else:
            print(f"[!] Warning: {md_file} not found.")

    print("\n[SUCCESS] All PDF, HTML, and DOCX files in Reports/ generated successfully.")
